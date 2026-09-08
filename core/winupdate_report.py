# -*- coding: utf-8 -*-
"""Windows Update 호환성 검증 체크리스트 xlsx에 판정 결과를 기록한다.

기준 문서: `..\\..\\Windows Update 호환성 검증 Checklist_Bellalun Viewer_R-25-782.xlsx`
시트 `Checklist` (헤더 5행, TC 6~18행, 기존 Result 열 K~R — 최신이 왼쪽).

`core/checklist.py`(기본기능 체크리스트)와 형식이 다르다 — 기본기능은 오른쪽에
새 열 묶음을 append하지만, 이 문서는 **기존 Result 열과 같은 형식으로 K열에
새 열을 삽입**해야 한다(`Result\\n<MM/DD>` 헤더 + `Pass`/`Fail`/`Manual` 값).
그래서 `core/checklist.write_results()`를 재사용하지 않고 이 모듈을 따로 둔다.
원본은 수정하지 않는다(사본에만 기록).
"""

import os
import shutil
from datetime import datetime

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

CHECKLIST_SHEET = "Checklist"
TC_ID_HEADER = "TC ID"
HEADER_ROW = 5
FIRST_TC_ROW = 6
LAST_TC_ROW = 18
NEW_RESULT_COL = 11  # K열

# 상단 실측값을 적는 행 (A열 라벨과 대응)
ENV_ROWS = {"os": 1, "os_version": 2, "os_build": 3, "viewer_version": 4}

FILLS = {
    "Pass":   PatternFill("solid", fgColor="C6EFCE"),
    "Fail":   PatternFill("solid", fgColor="FFC7CE"),
    "Manual": PatternFill("solid", fgColor="FFEB9C"),
}
FONTS = {
    "Pass":   Font(color="006100", bold=True),
    "Fail":   Font(color="9C0006", bold=True),
    "Manual": Font(color="9C6500", bold=True),
}

_VERDICT_LABEL = {"PASS": "Pass", "FAIL": "Fail", "MANUAL": "Manual",
                   "SKIP": "Skip", "BLOCKED": "Blocked"}

#: 자동화 전용 열 — 기존 Comment 열(삽입 후 밀린 위치) 오른쪽에 덧붙인다.
AUTOMATION_HEADERS = ["자동화 판정 일시", "재사용 근거 TC", "확인 항목 수", "자동화 note"]


def verdict_label(verdict):
    return _VERDICT_LABEL.get(str(verdict).upper(), str(verdict))


def _find_header_row(ws):
    for row in range(1, min(ws.max_row, 20) + 1):
        for col in range(1, ws.max_column + 1):
            if str(ws.cell(row, col).value or "").strip() == TC_ID_HEADER:
                return row, col
    raise ValueError(f"'{TC_ID_HEADER}' 헤더를 찾지 못했습니다.")


def write_results(source_xlsx, results, env=None, out_path=None, sheet_name=None,
                  issues=None):
    """WU 판정을 체크리스트 사본에 기록한다.

    results: `TCResult` 리스트(`tc_id`가 `TC_WindowsUpdate_01`.. 형식). 각
      결과 객체에 `wu_basis`(재사용한 TC 목록 문자열), `wu_known_issues`
      (아래 `issues`로 별도로도 모은다) 속성이 있으면 자동화 전용 열에 함께 적는다.
    env: {"os":.., "os_version":.., "os_build":.., "viewer_version":..} 상단
      1~4행에 실측값으로 기재한다.
    issues: [{"tc_id":, "kind":, "detail":, "impact":, "basis":}] — 별도
      'Issues' 시트에 기록한다. `kind`는 "기존 알려진 제품 결함" 또는
      "신규 발견"을 구분해서 넣는다.
    """
    if not os.path.isfile(source_xlsx):
        raise FileNotFoundError(source_xlsx)

    out_path = out_path or os.path.join(
        os.path.dirname(source_xlsx) or ".",
        f"WindowsUpdate_Checklist_Result_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    shutil.copyfile(source_xlsx, out_path)

    wb = load_workbook(out_path)
    ws = wb[sheet_name] if sheet_name else (
        wb[CHECKLIST_SHEET] if CHECKLIST_SHEET in wb.sheetnames else wb.active)
    hdr_row, tc_col = _find_header_row(ws)

    # 새 Result 열을 K에 **삽입**한다(기존 K~S가 오른쪽으로 밀린다). 이 시트는
    # 병합 셀이 없다고 실측 확인됐다(요청 프롬프트 기준 좌표표) — insert_cols가
    # 안전하다.
    ws.insert_cols(NEW_RESULT_COL)
    stamp_md = datetime.now().strftime("%m/%d")
    header_cell = ws.cell(hdr_row, NEW_RESULT_COL, f"Result\n{stamp_md}")
    header_cell.font = Font(bold=True)
    header_cell.alignment = Alignment(horizontal="center", vertical="center",
                                      wrap_text=True)
    ws.column_dimensions[header_cell.column_letter].width = 12

    env = env or {}
    for key, row in ENV_ROWS.items():
        value = env.get(key)
        if value:
            cell = ws.cell(row, NEW_RESULT_COL, value)
            cell.alignment = Alignment(horizontal="center", vertical="center")

    by_id = {r.tc_id: r for r in results}
    stamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    # 자동화 전용 열 — Comment 열(삽입으로 밀린 위치) 오른쪽부터.
    comment_col = None
    for col in range(1, ws.max_column + 1):
        if str(ws.cell(hdr_row, col).value or "").strip() == "Comment":
            comment_col = col
            break
    first_auto_col = (comment_col + 1) if comment_col else (ws.max_column + 1)
    auto_col_of = {}
    for i, name in enumerate(AUTOMATION_HEADERS):
        col = first_auto_col + i
        auto_col_of[name] = col
        cell = ws.cell(hdr_row, col, name)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = 30

    written = []
    for row in range(FIRST_TC_ROW, LAST_TC_ROW + 1):
        tc_id = str(ws.cell(row, tc_col).value or "").strip()
        if not tc_id:
            continue
        result = by_id.get(tc_id)
        if result is None:
            continue
        label = verdict_label(result.verdict)
        cell = ws.cell(row, NEW_RESULT_COL, label)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if label in FILLS:
            cell.fill = FILLS[label]
            cell.font = FONTS[label]
        ws.cell(row, auto_col_of["자동화 판정 일시"], stamp)
        basis = getattr(result, "wu_basis", "") or ""
        basis_cell = ws.cell(row, auto_col_of["재사용 근거 TC"], basis)
        basis_cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row, auto_col_of["확인 항목 수"], len(result.checks))
        note_lines = [c.note for c in result.checks if c.note]
        note_cell = ws.cell(row, auto_col_of["자동화 note"], "\n".join(note_lines))
        note_cell.alignment = Alignment(vertical="top", wrap_text=True)
        written.append(tc_id)

    unmatched = sorted(set(by_id) - set(written))

    # --- Issues 시트 ---------------------------------------------------
    issues = list(issues or [])
    for result in results:
        issues.extend(getattr(result, "wu_known_issues", None) or [])
    if issues:
        if "Issues" in wb.sheetnames:
            del wb["Issues"]
        iws = wb.create_sheet("Issues")
        headers = ["TC ID", "구분", "내용", "WU 판정 영향", "근거"]
        for col, name in enumerate(headers, 1):
            cell = iws.cell(1, col, name)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        widths = [22, 20, 50, 30, 40]
        for col, width in enumerate(widths, 1):
            iws.column_dimensions[iws.cell(1, col).column_letter].width = width
        for i, item in enumerate(issues, 2):
            iws.cell(i, 1, item.get("tc_id", ""))
            iws.cell(i, 2, item.get("kind", ""))
            c = iws.cell(i, 3, item.get("detail", ""))
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c = iws.cell(i, 4, item.get("impact", ""))
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c = iws.cell(i, 5, item.get("basis", ""))
            c.alignment = Alignment(vertical="top", wrap_text=True)

    wb.save(out_path)
    return {"path": out_path, "written": written, "unmatched": unmatched,
            "issues": len(issues), "sheet": ws.title}
