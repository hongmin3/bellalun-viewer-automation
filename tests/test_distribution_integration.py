"""Additional safe regression tests to ship with the engine integration."""
from contextlib import contextmanager, nullcontext, redirect_stdout
import io
import json
import os
from pathlib import Path
import subprocess
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from bellalunauto import app
from bellalunauto.runtime import AlreadyRunning, Layout
from core import specs, winupdate_report
import run


class EngineIntegrationExtraTests(unittest.TestCase):
    def test_context_routes_explicit_ocr_path_to_already_imported_helpers(self):
        from core import screen, viewer_processing
        with tempfile.TemporaryDirectory(prefix='OCR 한글 ') as folder:
            root = Path(folder)
            config = root/'config.json'
            selected = r'D:\OCR 한글 경로\tesseract.exe'
            config.write_text(json.dumps({'data_dir':str(root),
                                          'xipl':{'tesseract_exe':selected}}), encoding='utf-8')
            with patch.dict(os.environ, {'BELLALUN_TESSERACT':r'E:\inherited\tesseract.exe'}), \
                 patch.object(screen, 'TESSERACT_EXE', 'already imported screen default'), \
                 patch.object(viewer_processing, 'TESSERACT_EXE', 'already imported processing default'), \
                 patch.object(screen, 'ocr', side_effect=AssertionError('OCR execution')), \
                 patch.object(screen, 'grab', side_effect=AssertionError('screen capture')), \
                 patch.object(subprocess, 'run', side_effect=AssertionError('process execution')), \
                 patch.object(run, 'BellalunDb', return_value=SimpleNamespace()):
                run.Context(str(config))
                self.assertEqual(screen.TESSERACT_EXE, selected)
                self.assertEqual(viewer_processing.TESSERACT_EXE, selected)

    def test_context_missing_null_or_empty_ocr_keeps_existing_helper_defaults(self):
        from core import screen, viewer_processing
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            config = root/'config.json'
            for optional in ({}, {'xipl':None}, {'xipl':{}}, {'xipl':{'tesseract_exe':''}}):
                config.write_text(json.dumps({'data_dir':str(root), **optional}), encoding='utf-8')
                inherited = r'E:\custom environment\tesseract.exe'
                with self.subTest(optional=optional), \
                     patch.dict(os.environ, {'BELLALUN_TESSERACT':inherited}), \
                     patch.object(screen, 'TESSERACT_EXE', inherited), \
                     patch.object(viewer_processing, 'TESSERACT_EXE', inherited), \
                     patch.object(screen, 'ocr', side_effect=AssertionError('OCR execution')), \
                     patch.object(screen, 'grab', side_effect=AssertionError('screen capture')), \
                     patch.object(subprocess, 'run', side_effect=AssertionError('process execution')), \
                     patch.object(run, 'BellalunDb', return_value=SimpleNamespace()):
                    run.Context(str(config))
                    self.assertEqual(screen.TESSERACT_EXE, inherited)
                    self.assertEqual(viewer_processing.TESSERACT_EXE, inherited)

    def test_engine_exports_workspace_resources_and_replay_without_worker_flag(self):
        with tempfile.TemporaryDirectory(prefix='engine 한글 ') as folder:
            root = Path(folder)
            layout = Layout(root/'설치 자료', root/'작업 결과', root/'설정 파일.json')
            observed = {}
            def engine_main():
                observed.update({name: os.environ.get(name) for name in
                    ('BELLALUN_WORKSPACE', 'BELLALUN_RESOURCE_DIR', 'BELLALUN_LAUNCH_COMMAND')})
                return 19
            executable = 'E:\\한글 설치\\BellalunAuto.exe'
            with patch.dict(os.environ), patch.object(run, 'main', side_effect=engine_main), \
                 patch.object(app, 'child_command', side_effect=lambda args: [executable] + args):
                self.assertEqual(app._engine(layout, ['list']), 19)
            self.assertEqual(observed['BELLALUN_WORKSPACE'], str(layout.workspace))
            self.assertEqual(observed['BELLALUN_RESOURCE_DIR'], str(layout.resource_dir))
            self.assertEqual(observed['BELLALUN_LAUNCH_COMMAND'], subprocess.list2cmdline([
                executable, '--workspace', str(layout.workspace), '--config', str(layout.config_path), 'list']))
            self.assertNotIn('--worker', observed['BELLALUN_LAUNCH_COMMAND'])

    def test_guard_records_own_pid_as_child_lock_owner(self):
        from core import automation_health, ui
        from bellalunauto import processes
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            config = root/'config.json'
            config.write_text('{}', encoding='utf-8')
            child = SimpleNamespace(pid=123, stdout=io.StringIO(''), poll=lambda:0, returncode=0)
            with patch.object(app, 'RunLock', return_value=nullcontext()), \
                 patch.object(ui, 'foreground_unlocked', return_value=nullcontext()), \
                 patch.object(ui, 'taskbar_autohidden', return_value=nullcontext()), \
                 patch.object(processes, 'WindowsJob'), \
                 patch.object(app.subprocess, 'Popen', return_value=child) as start, \
                 patch.object(app, 'completed_report', return_value={'verdicts':{}}), \
                 patch.object(automation_health, 'write_state'), \
                 patch.object(automation_health, 'notify_windows'):
                self.assertEqual(app.run_guarded(Layout(root, root, config), ['run-wf01']), 0)
            self.assertEqual(start.call_args.kwargs['env']['BELLALUN_PRODUCT_LOCK_OWNER'], str(os.getpid()))

    def test_raw_main_takes_legacy_lock_before_dispatch(self):
        self.assertTrue(callable(getattr(run, '_main', None)))
        with patch.object(app, 'legacy_product_lock', return_value=nullcontext()) as lock, \
             patch.object(run, '_main', return_value=21), \
             patch.object(run.sys, 'argv', ['run.py', '--config', 'config.json', 'list']):
            self.assertEqual(run.main(), 21)
        lock.assert_called_once_with(['--config', 'config.json', 'list'])

    def test_raw_help_and_readonly_commands_never_touch_ui_settings(self):
        from core import ui
        for arguments in (['--help'], ['list'], ['run-regression', '--help'],
                          ['run-winupdate', '--from-regression=old.json']):
            with self.subTest(arguments=arguments), \
                 patch.object(app, 'RunLock', side_effect=AssertionError('readonly lock')), \
                 patch.object(ui, 'foreground_unlocked', side_effect=AssertionError('foreground mutation')), \
                 patch.object(ui, 'taskbar_autohidden', side_effect=AssertionError('taskbar mutation')), \
                 patch.object(run, '_main', return_value=0), \
                 patch.object(run.sys, 'argv', ['run.py'] + arguments):
                self.assertEqual(run.main(), 0)

    def test_busy_raw_source_rejects_before_ui_settings(self):
        from core import ui
        with patch.object(app, 'legacy_product_lock', side_effect=AlreadyRunning('busy')), \
             patch.object(ui, 'foreground_unlocked', side_effect=AssertionError('foreground mutation')), \
             patch.object(ui, 'taskbar_autohidden', side_effect=AssertionError('taskbar mutation')), \
             patch.object(run, '_main', side_effect=AssertionError('engine execution')), \
             patch.object(run.sys, 'argv', ['run.py', 'run-regression']), \
             redirect_stdout(io.StringIO()):
            self.assertEqual(run.main(), 4)

    def test_product_ui_contexts_begin_only_after_source_lock(self):
        from core import ui
        events = []
        @contextmanager
        def record(name):
            events.append(name)
            yield
        def engine():
            events.append('engine')
            return 23
        with patch.object(app, 'legacy_product_lock', side_effect=lambda args: record('lock')), \
             patch.object(ui, 'foreground_unlocked', side_effect=lambda: record('foreground')), \
             patch.object(ui, 'taskbar_autohidden', side_effect=lambda: record('taskbar')), \
             patch.object(run, '_main', side_effect=engine), \
             patch.object(run.sys, 'argv', ['run.py', 'run-regression']):
            self.assertEqual(run.main(), 23)
        self.assertEqual(events, ['lock', 'foreground', 'taskbar', 'engine'])

    def test_wu_mode_and_abort_state_are_visible_in_console_and_report_environment(self):
        with tempfile.TemporaryDirectory() as folder:
            ctx = SimpleNamespace(reports_root=folder)
            for mode, completed, expected_mode, expected_state in (
                    ('historical', False, '과거 결과 변환', '해당 없음'),
                    ('live', False, '실시간 제품 시험', '조기 중단'),
                    ('live', True, '실시간 제품 시험', '완료')):
                output = io.StringIO()
                with self.subTest(mode=mode, completed=completed), \
                     patch.object(run, '_winupdate_source_path', return_value='fixture.xlsx'), \
                     patch.object(run, '_winupdate_env', return_value={'os':'fixture OS'}), \
                     patch.object(winupdate_report, 'write_results', return_value={
                         'path':'output.xlsx', 'written':[], 'unmatched':[], 'issues':0}) as xlsx, \
                     patch.object(run, 'write_reports', return_value={}) as report, \
                     patch.object(run, 'announce_done'), redirect_stdout(output):
                    run._finish_winupdate(ctx, [], execution_mode=mode,
                                         chain_completed=completed, source_note='fixture source')
                env = report.call_args.kwargs['meta']['env']
                self.assertEqual(env['실행 방식'], expected_mode)
                self.assertEqual(env['체인 완료'], expected_state)
                self.assertEqual(env['근거'], 'fixture source')
                self.assertEqual(xlsx.call_args.kwargs['env'], env)
                self.assertIn(expected_mode, output.getvalue())
                self.assertIn(expected_state, output.getvalue())

    def test_xlsx_preserves_visible_execution_mode_and_original_workbook(self):
        from openpyxl import Workbook, load_workbook
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source, output = root/'source.xlsx', root/'output.xlsx'
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = 'Checklist'
            sheet.cell(5, 1, 'TC ID')
            sheet.cell(6, 1, 'TC_WindowsUpdate_01')
            workbook.save(source)
            original = source.read_bytes()
            env = {'실행 방식':'과거 결과 변환', '체인 완료':'해당 없음', '근거':'fixture source'}
            winupdate_report.write_results(str(source), [], env=env, out_path=str(output))
            saved = load_workbook(output)
            self.assertIn('실행 정보', saved.sheetnames)
            visible = dict(saved['실행 정보'].iter_rows(min_row=2, max_col=2, values_only=True))
            self.assertEqual(visible, env)
            self.assertEqual(source.read_bytes(), original)

    def test_pdf_content_change_invalidates_cache_even_with_same_mtime(self):
        from pypdf import PdfWriter
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            pdf = root/'사양서1.pdf'
            cache = root/'workspace/Cache/Specs'
            writer = PdfWriter()
            writer.add_blank_page(width=72, height=72)
            writer.write(pdf)
            before = pdf.stat()
            self.assertEqual(len(specs.extract(str(pdf), cache_dir=str(cache))), 1)
            writer.add_blank_page(width=72, height=72)
            writer.write(pdf)
            os.utime(pdf, ns=(before.st_atime_ns, before.st_mtime_ns))
            self.assertEqual(len(specs.extract(str(pdf), cache_dir=str(cache))), 2)
            self.assertEqual(len(list(cache.glob('*.txt'))), 2)
            self.assertEqual(list(root.glob('*.txt')), [])


if __name__ == '__main__':
    unittest.main()
